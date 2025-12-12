import os
import re
import PyPDF2
from docx import Document


class DocumentProcessor:
    def __init__(self):
        self.target_phrases = [r"\bсроки\b", r"\bсрок\b", r"\bдата\b"]

        self.date_patterns = [
            r"\b(\d{2}\.\d{2}\.\d{4})(?=(?: |, |\. ))",
            r"\b(\d{2}\.\d{2}\.\d{2})(?=(?: |, |\. ))",
            r"\b(\d{4}-\d{2}-\d{2})(?=(?: |, |\. ))",
            r"\b(\d{2}/\d{2}/\d{4})(?=(?: |, |\. ))",
            r"\b(\d{2}/\d{2}/\d{2})(?=(?: |, |\. ))",
            r"(?:с\s+)(\d{1,2}\.\d{1,2}\.\d{4})(?:\s+по\s+)"
            r"(\d{1,2}\.\d{1,2}\.\d{4})",
            r"(?:до\s+)(\d{1,2}\.\d{1,2}\.\d{4})",
            r"(?:по\s+)(\d{1,2}\.\d{1,2}\.\d{4})",
        ]

        self.city_pattern = re.compile(
            r"(г\. [А-ЯЁ][а-яё]{2,}(?:[ -][А-ЯЁ][а-яё]{2,})*)"
        )

    def find_dates(self, text):
        if not text:
            return []
        dates = []
        for pattern in self.date_patterns:
            for match in re.finditer(pattern, text, re.IGNORECASE):
                if len(match.groups()) == 2:
                    dates.extend([match.group(1), match.group(2)])
                else:
                    dates.append(match.group(1))
        return list(set(dates))

    def find_cities(self, text):
        if not text:
            return []
        cities = set()
        for m in self.city_pattern.finditer(text):
            cities.add(m.group(1))
        return list(cities)

    def process_file(self, file_path, mode="date"):
        if file_path.lower().endswith(".docx"):
            text = self.docx_handler(file_path)
        elif file_path.lower().endswith(".pdf"):
            text = self.pdf_handler(file_path)
        else:
            print("Unsupported file format")
            return

        if not text:
            print("Failed to extract text")
            return

        search_dates = mode in ("date", "both")
        search_cities = mode in ("city", "both")

        dates = []
        cities = []

        if search_dates:
            dates = self.find_dates(text)
        if search_cities:
            cities = self.find_cities(text)

        if search_dates:
            if dates:
                print("Dates:")
                for date in dates:
                    print(date)
            else:
                print("No dates found")

        if search_cities:
            if cities:
                print("Cities:")
                for city in cities:
                    print(city)
            else:
                print("No cities found in format 'г. <City>'")

    def docx_handler(self, file_path):
        try:
            doc = Document(file_path)
            parts = []

            for paragraph in doc.paragraphs:
                if paragraph.text:
                    parts.append(paragraph.text)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            parts.append(cell.text)

            return "\n".join(parts)
        except Exception as e:
            print(f"Error reading DOCX file {file_path}: {e}")
            return ""

    def pdf_handler(self, file_path):
        try:
            text = ""
            with open(file_path, "rb") as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            return text
        except Exception as e:
            print(f"Error reading PDF file {file_path}: {e}")
            return ""

    def extract(self, file_path, mode="both"):
        if file_path.lower().endswith(".docx"):
            text = self.docx_handler(file_path)
        elif file_path.lower().endswith(".pdf"):
            text = self.pdf_handler(file_path)
        else:
            return [], []

        if not text:
            return [], []

        dates = []
        cities = []
        if mode in ("date", "both"):
            dates = self.find_dates(text)
        if mode in ("city", "both"):
            cities = self.find_cities(text)
        return dates, cities


def main():
    processor = DocumentProcessor()

    paths_str = input("Введите путь(и) к файлу(ам), через пробел: ").strip()
    if not paths_str:
        print("Путь не указан")
        return

    paths = paths_str.split()

    with open("data_city.txt", "a", encoding="utf-8") as out:
        for file_path in paths:
            if not os.path.exists(file_path):
                print(f"File not found: {file_path}")
                continue

            dates, cities = processor.extract(file_path, mode="both")

            filename = os.path.basename(file_path)
            dates_str = ", ".join(dates) if dates else "-"
            cities_str = ", ".join(cities) if cities else "-"

            line = f"{filename} | {dates_str} | {cities_str}\n"
            out.write(line)

    print("Результат записан в data_city.txt")


if __name__ == "__main__":
    main()
